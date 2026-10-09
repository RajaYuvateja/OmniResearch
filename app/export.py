import io, re, html

def to_pdf(md: str, chart: bytes | None) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0e6b5c"),
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        "CustomH2",
        parent=styles["Heading2"],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#1e293b"),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "CustomBody",
        parent=styles["BodyText"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )
    
    buf = io.BytesIO()
    flow = []

    def safe_p(text: str, style):
        # Escape HTML / XML special characters
        clean = html.escape(text)
        # Re-enable safe simple inline tags if requested
        clean = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", clean)
        clean = re.sub(r"\*(.*?)\*", r"<i>\1</i>", clean)
        try:
            return Paragraph(clean, style)
        except Exception:
            # Fallback to absolute plain text without tags if ReportLab XML parser complains
            return Paragraph(html.escape(text), style)

    for ln in md.splitlines():
        trimmed = ln.strip()
        if not trimmed:
            flow.append(Spacer(1, 4))
            continue
        m = re.match(r"(#+)\s+(.*)", trimmed)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2)
            style = title_style if level == 1 else h2_style
            flow.append(safe_p(heading_text, style))
            if heading_text.lower().startswith("analysis") and chart:
                flow.append(Spacer(1, 6))
                try:
                    flow.append(Image(io.BytesIO(chart), width=440, height=226))
                    flow.append(Spacer(1, 6))
                except Exception:
                    pass
        else:
            flow.append(safe_p(trimmed, body_style))

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    doc.build(flow)
    return buf.getvalue()

def to_docx(md: str, chart: bytes | None) -> bytes:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    d = Document()
    
    for ln in md.splitlines():
        trimmed = ln.strip()
        m = re.match(r"(#+)\s+(.*)", trimmed)
        if m:
            level = min(len(m.group(1)), 3)
            h = d.add_heading(m.group(2), level=level)
            if m.group(2).lower().startswith("analysis") and chart:
                try:
                    d.add_picture(io.BytesIO(chart), width=Inches(5.8))
                except Exception:
                    pass
        elif trimmed:
            is_bullet = trimmed.startswith(("-", "*", "•"))
            text = trimmed.lstrip("-*• ")
            d.add_paragraph(text, style="List Bullet" if is_bullet else None)
            
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

def to_pptx(md: str, chart: bytes | None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    p = Presentation()
    sections, cur = [], None
    
    for ln in md.splitlines():
        m = re.match(r"#+\s+(.*)", ln.strip())
        if m:
            cur = [m.group(1), []]
            sections.append(cur)
        elif cur and ln.strip():
            cur[1].append(ln.strip().lstrip("-*• ")[:180])
            
    for title, lines in sections[:12]:
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = title[:80]
        tf = s.placeholders[1].text_frame
        tf.clear()
        for i, t in enumerate(lines[:5]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = t
            para.font.size = Pt(15)
            
    if chart:
        try:
            s = p.slides.add_slide(p.slide_layouts[5])
            s.shapes.title.text = "Market & Empirical Analytics"
            s.shapes.add_picture(io.BytesIO(chart), Inches(1), Inches(1.5), width=Inches(8))
        except Exception:
            pass
            
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()
