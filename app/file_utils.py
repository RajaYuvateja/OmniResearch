"""Utility for extracting and processing text from uploaded documents (PDF, DOCX, TXT, CSV, JSON, MD)."""
import io
import os
import json
import csv
from typing import Dict, Any, Optional

def extract_text_from_file(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses uploaded file bytes into extracted text based on file extension.
    Returns dictionary with text, filename, char_count, and file_type.
    """
    ext = os.path.splitext(filename)[1].lower().strip(".")
    text = ""
    file_type = ext.upper() if ext else "UNKNOWN"

    try:
        if ext == "pdf":
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for i, page in enumerate(reader.pages):
                page_content = page.extract_text() or ""
                if page_content.strip():
                    pages_text.append(f"--- [Page {i + 1}] ---\n{page_content.strip()}")
            text = "\n\n".join(pages_text)

        elif ext in ("docx", "doc"):
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)

        elif ext in ("txt", "md", "markdown", "rst", "log"):
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="replace")

        elif ext in ("csv", "tsv"):
            delimiter = "\t" if ext == "tsv" else ","
            try:
                decoded = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                decoded = file_bytes.decode("latin-1", errors="replace")
            
            reader = csv.reader(io.StringIO(decoded), delimiter=delimiter)
            rows = []
            for row in reader:
                if any(row):
                    rows.append(" | ".join(str(cell).strip() for cell in row))
            text = "\n".join(rows[:1000])  # limit to first 1000 rows

        elif ext == "json":
            try:
                decoded = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                decoded = file_bytes.decode("latin-1", errors="replace")
            parsed = json.loads(decoded)
            text = json.dumps(parsed, indent=2)

        else:
            # Fallback to plain text attempt
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="replace")

    except Exception as e:
        text = f"Error extracting text from {filename}: {str(e)}"

    # Normalize whitespace
    text = text.strip()
    
    # Cap excessive size to around 80,000 characters to prevent LLM context saturation
    max_chars = 80000
    is_truncated = False
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n\n[... Truncated after {max_chars} characters for optimal processing ...]"
        is_truncated = True

    return {
        "text": text,
        "filename": filename,
        "size_bytes": len(file_bytes),
        "char_count": len(text),
        "file_type": file_type,
        "is_truncated": is_truncated
    }

def chunk_text(text: str, chunk_size: int = 1500, overlap: int = 200) -> list[str]:
    """Splits a document text into readable overlapping chunks for retrieval."""
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks
